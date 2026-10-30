"""文件 → 纯文本。

全链路只用宽松许可（MIT / BSD / Apache-2.0）的依赖，不含任何 AGPL 组件：

- 图片：Pillow 解码 + EXIF 纠正 → RapidOCR（Apache-2.0）识别
- PDF 文本层：pypdf（BSD-3-Clause）抽取；若抽出来是乱码（中文 CID 字体常见），
  自动改用 pypdfium2（Apache-2.0）再抽一次
- PDF 扫描件：pypdfium2 把页面渲染成位图 → 交给同一套 OCR 引擎
- Word：python-docx；纯文本：标准库直读

对上层只暴露 extract() 和 ExtractError。
"""

from __future__ import annotations

import io
import logging
import re
import threading
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

log = logging.getLogger("extract")

Progress = Callable[[int, str], None]

# 支持的文件类型
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
DOC_EXTS = {".pdf", ".docx"}
TEXT_EXTS = {".txt", ".md"}
ALLOWED_EXTS = IMAGE_EXTS | DOC_EXTS | TEXT_EXTS

# 文件头魔数：拦住"扩展名对、内容不对"的假文件
SIGNATURES = {
    ".pdf": (b"%PDF", "PDF"),
    ".docx": (b"PK", "DOCX"),
    ".png": (b"\x89PNG", "PNG"),
    ".jpg": (b"\xff\xd8", "JPEG"),
    ".jpeg": (b"\xff\xd8", "JPEG"),
    ".tif": (b"II*\x00", "TIFF"),
    ".tiff": (b"II*\x00", "TIFF"),
}

OCR_DPI = 200                 # 再高对识别率提升有限，耗时翻倍
SCAN_CHARS_PER_PAGE = 20      # 平均每页少于这么多字 → 判定为扫描件
MAX_PAGES = 200               # 文档页数硬上限
OCR_PAGE_LIMIT = 50           # 扫描件最多 OCR 这么多页，其余跳过

# 正常中英文里几乎不会出现的字符（Latin-1 补充与扩展区）
_SUSPICIOUS = re.compile(r"[\u0080-\u024f]")


class ExtractError(Exception):
    """抽取阶段的失败。kind 供上层映射 HTTP 状态码。"""

    def __init__(self, kind: str, message: str) -> None:
        super().__init__(message)
        self.kind = kind
        self.message = message


@dataclass
class Extracted:
    text: str
    kind: str                 # text / docx / image / pdf-text / pdf-ocr
    pages: int = 0
    used_ocr: bool = False

    @property
    def chars(self) -> int:
        return len(re.sub(r"\s+", "", self.text))


def sniff(head: bytes, ext: str) -> None:
    """上传落盘后立刻校验文件头，早失败比晚失败好。"""
    sig = SIGNATURES.get(ext)
    if sig and head and not head.startswith(sig[0]):
        raise ExtractError("unsupported", f"文件内容不像 {sig[1]}，请确认扩展名是否正确")


def extract(path: Path, on_progress: Progress | None = None) -> Extracted:
    """按扩展名分发。任何失败都抛 ExtractError，不泄露底层异常类型给前端。"""
    if not path.is_file():
        raise ExtractError("bad_file", "文件不存在或已被清理")

    ext = path.suffix.lower()
    if ext not in ALLOWED_EXTS:
        raise ExtractError("unsupported", f"不支持 {ext or '无扩展名'}；可用类型：{' '.join(sorted(ALLOWED_EXTS))}")

    if ext in TEXT_EXTS:
        return Extracted(_tidy(_read_plain(path)), "text")
    if ext == ".docx":
        return Extracted(_tidy(_read_docx(path)), "docx")
    if ext in IMAGE_EXTS:
        if on_progress:
            on_progress(20, "正在识别图片文字")
        return Extracted(_tidy(_read_image(path)), "image", pages=1, used_ocr=True)
    return _read_pdf(path, on_progress)


# ------------------------------ 纯文本 ------------------------------

def _read_plain(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractError("bad_file", "文件编码无法识别，请另存为 UTF-8 后重试")


def _tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u3000]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


# ------------------------------ Word ------------------------------

def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ExtractError("missing_dep", "缺少 python-docx，请重新运行 run.bat") from exc

    try:
        doc = Document(str(path))
    except Exception as exc:
        raise ExtractError("bad_file", f"这个 .docx 打不开：{exc}") from exc

    lines = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]

    # 表格不能丢：按行拼回文本，单元格之间用 | 分隔
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


# ------------------------------ 图片 ------------------------------

def _read_image(path: Path) -> str:
    try:
        from PIL import Image, ImageOps
    except ImportError as exc:
        raise ExtractError("missing_dep", "缺少 pillow，请重新运行 run.bat") from exc

    try:
        with Image.open(str(path)) as img:
            img = ImageOps.exif_transpose(img)   # 手机竖拍不纠正，OCR 出来全是乱码
            img = img.convert("RGB")
            buffer = io.BytesIO()
            img.save(buffer, format="PNG")
            png = buffer.getvalue()
    except Exception as exc:
        raise ExtractError("bad_file", f"这张图打不开：{exc}") from exc
    return "\n".join(ocr(png))


# ------------------------------ PDF ------------------------------

def _read_pdf(path: Path, on_progress: Progress | None) -> Extracted:
    """先找文本层（快），有且正常就直接用；否则渲染成图走 OCR（慢但通用）。"""
    pages, text = _pdf_probe_text(path)

    if pages == 0:
        raise ExtractError("bad_file", "PDF 里没有页面")
    if pages > MAX_PAGES:
        raise ExtractError("unsupported", f"超过 {MAX_PAGES} 页，请拆分后再上传")

    if len(re.sub(r"\s+", "", text)) >= SCAN_CHARS_PER_PAGE * pages:
        return Extracted(_tidy(text), "pdf-text", pages=pages)

    if on_progress:
        on_progress(15, f"没有可用文本层（{pages} 页），改用 OCR")

    lines = _pdf_ocr(path, pages, on_progress)
    if pages > OCR_PAGE_LIMIT:
        lines.append(f"（本 PDF 共 {pages} 页，仅识别前 {OCR_PAGE_LIMIT} 页）")
    return Extracted(_tidy("\n".join(lines)), "pdf-ocr", pages=pages, used_ocr=True)


def _pdf_probe_text(path: Path) -> tuple[int, str]:
    """先 pypdf，抽出来是乱码就换 pypdfium2，都不行就交给 OCR。"""
    pages, text = 0, ""
    try:
        pages, text = _pdf_text_pypdf(path)
    except Exception as exc:                      # 打不开 / 加密 / 解析异常
        log.warning("pypdf 抽取失败：%s", exc)

    if _text_is_garbage(text):
        try:
            alt = _pdf_text_pdfium(path)
        except Exception as exc:
            log.warning("pypdfium2 抽取文本失败：%s", exc)
            alt = ""
        if not _text_is_garbage(alt):
            log.info("pypdf 结果不可用，已改用 pypdfium2 抽取文本层")
            text = alt

    if pages == 0:
        pages = _pdf_page_count(path)
    return pages, text


def _pdf_text_pypdf(path: Path) -> tuple[int, str]:
    """用 pypdf（BSD-3-Clause）抽文本层。"""
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractError("missing_dep", "缺少 pypdf，请重新运行 run.bat") from exc

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise ExtractError("bad_file", f"这个 PDF 打不开：{exc}") from exc

    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")                    # 很多 PDF 只设了权限密码
        except Exception as exc:
            raise ExtractError("bad_file", "这个 PDF 有密码保护，请先解除加密再上传") from exc

    pages = len(reader.pages)
    chunks: list[str] = []
    for page in reader.pages:
        try:
            chunks.append(page.extract_text() or "")
        except Exception:                         # 单页失败不影响整份
            chunks.append("")
    return pages, "\n".join(chunks)


def _pdf_text_pdfium(path: Path) -> str:
    """用 pypdfium2（Apache-2.0）抽文本层，对中文 CID 字体更可靠。"""
    pdfium = _import_pdfium()
    pdf = pdfium.PdfDocument(str(path))
    parts: list[str] = []
    try:
        for index in range(len(pdf)):
            page = pdf[index]
            try:
                textpage = page.get_textpage()
                try:
                    read = getattr(textpage, "get_text_bounded", None) or textpage.get_text_range
                    parts.append(read() or "")
                finally:
                    _safe_close(textpage)
            finally:
                _safe_close(page)
    finally:
        _safe_close(pdf)
    return "\n".join(parts)


def _pdf_page_count(path: Path) -> int:
    try:
        pdfium = _import_pdfium()
        pdf = pdfium.PdfDocument(str(path))
        try:
            return len(pdf)
        finally:
            _safe_close(pdf)
    except Exception as exc:
        raise ExtractError("bad_file", f"这个 PDF 打不开：{exc}") from exc


def _pdf_ocr(path: Path, pages: int, on_progress: Progress | None) -> list[str]:
    """把 PDF 页面渲染成位图，逐页交给 OCR。"""
    pdfium = _import_pdfium()
    try:
        pdf = pdfium.PdfDocument(str(path))
    except Exception as exc:
        raise ExtractError("bad_file", f"这个 PDF 打不开：{exc}") from exc

    todo = min(pages, OCR_PAGE_LIMIT)
    lines: list[str] = []
    try:
        for index in range(todo):
            lines.extend(ocr(_render_page(pdf, index)))
            if on_progress and (index + 1 == todo or (index + 1) % 3 == 0):
                on_progress(15 + int(50 * (index + 1) / todo), f"OCR 第 {index + 1}/{todo} 页")
    finally:
        _safe_close(pdf)
    return lines


def _render_page(pdf: Any, index: int) -> bytes:
    """把某一页渲染成 PNG 字节（pypdfium2 → PIL → BytesIO）。"""
    page = pdf[index]
    bitmap = None
    try:
        bitmap = page.render(scale=OCR_DPI / 72.0)    # PDF 坐标系按 72 dpi，按比例放大
        image = bitmap.to_pil()
        if image.mode != "RGB":
            image = image.convert("RGB")
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        return buffer.getvalue()
    except Exception as exc:
        raise ExtractError("bad_file", f"第 {index + 1} 页渲染失败：{exc}") from exc
    finally:
        _safe_close(bitmap)
        _safe_close(page)


def _import_pdfium() -> Any:
    try:
        import pypdfium2
    except ImportError as exc:
        raise ExtractError("missing_dep", "缺少 pypdfium2，请重新运行 run.bat") from exc
    return pypdfium2


def _safe_close(obj: Any) -> None:
    """pypdfium2 的对象建议显式 close，但版本间有差异，这里做保护。"""
    if obj is None:
        return
    close = getattr(obj, "close", None)
    if callable(close):
        try:
            close()
        except Exception:
            pass


def _text_is_garbage(text: str) -> bool:
    """抽出来的文本是否不可用（空、过短、或明显是乱码）。"""
    clean = re.sub(r"\s+", "", text)
    if len(clean) < 8:
        return True
    return not _text_looks_natural(clean)


def _text_looks_natural(text: str) -> bool:
    """判断文本像不像正常文字。

    pypdf 遇到中文 CID 字体（UniGB-UTF16-H 之类）时，会把字节直接映射成拉丁
    字符，产出「fzýR6N§~¿[c^¦YvØb¥TJ」这种有长度、没意义的东西。
    正常中英文里几乎不会出现拉丁补充/扩展区字符，用它的占比就能区分开。
    """
    clean = re.sub(r"\s+", "", text)
    if len(clean) < 8:
        return False
    suspicious = len(_SUSPICIOUS.findall(clean))
    return suspicious / len(clean) <= 0.02


# ------------------------------ OCR ------------------------------

_ENGINE: Any = None
_ENGINE_LOCK = threading.Lock()


def _get_engine() -> Any:
    """OCR 模型只加载一次，多线程共用；第一次约 1-2 秒。"""
    global _ENGINE
    if _ENGINE is not None:
        return _ENGINE
    with _ENGINE_LOCK:
        if _ENGINE is None:
            try:
                from rapidocr_onnxruntime import RapidOCR
            except ImportError as exc:
                raise ExtractError("missing_dep", "缺少 rapidocr-onnxruntime，请重新运行 run.bat") from exc
            log.info("加载 OCR 模型（首次约 1-2 秒）")
            _ENGINE = RapidOCR()
    return _ENGINE


def ocr(source: bytes) -> list[str]:
    """识别一段图片字节，返回逐行文本。"""
    engine = _get_engine()
    try:
        result, _ = engine(source)
    except Exception as exc:
        raise ExtractError("bad_file", f"OCR 识别失败：{exc}") from exc

    if result is None:
        return []

    # 新版 rapidocr 返回带 .txts 的对象
    if hasattr(result, "txts"):
        return [str(t).strip() for t in result.txts if str(t).strip()]

    lines: list[str] = []
    for item in result:
        text = ""
        if isinstance(item, dict):
            text = str(item.get("txt") or item.get("text") or "")
        elif isinstance(item, (list, tuple)) and len(item) >= 2:
            text = str(item[1])
        text = text.strip()
        if text:
            lines.append(text)
    return lines
